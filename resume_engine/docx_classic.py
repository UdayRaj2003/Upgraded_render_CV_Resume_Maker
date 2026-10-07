"""Simple ATS Word + PDF matching Resume_Uday_26_08.docx."""

from __future__ import annotations

import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from ruamel.yaml import YAML

try:
    from docx import Document
    from docx.enum.text import WD_LINE_SPACING
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Inches, Pt, RGBColor
except ImportError:  # pragma: no cover
    Document = None  # type: ignore[misc, assignment]
    RGBColor = None  # type: ignore[misc, assignment]


_MONTHS = (
    "January",
    "February",
    "March",
    "April",
    "May",
    "June",
    "July",
    "August",
    "September",
    "October",
    "November",
    "December",
)
_LABEL = re.compile(r"^(\s*)(\*\*)?([^:*]+?)(\*\*)?\s*:\s*(.*)$", re.DOTALL)
_URL_CHUNK = re.compile(
    r"(https?://[^\s]+|"
    r"(?:www\.)?(?:drive\.google\.com|github\.com|linkedin\.com|leetcode\.com|"
    r"geeksforgeeks\.org|hackerrank\.com)[^\s]*)",
    re.I,
)

FONT = "Calibri"
SIZE_LABEL = 12.0
SIZE_BODY = 10.0
BLACK = RGBColor(0, 0, 0) if RGBColor is not None else None
LINK_BLUE = RGBColor(0, 0, 255) if RGBColor is not None else None
LINK_BLUE_HEX = "0000FF"
LINK_BLUE_RGB = (0, 0, 255)

# Tight page: ~0.05in top (was 0.10in plus Word header gap)
DOCX_TOP_IN = 0.05
DOCX_EDGE_IN = 0.08


@dataclass
class AtsRun:
    text: str
    bold: bool = False
    size: float = 10.0
    url: str | None = None


@dataclass
class AtsPara:
    runs: list[AtsRun] = field(default_factory=list)
    indent_cm: float = 0.0


def _plain(text: str) -> str:
    return re.sub(r"\*\*(.+?)\*\*", r"\1", str(text or ""))


def _strip_link(raw: str) -> str:
    text = (raw or "").strip()
    match = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", text)
    return match.group(1) if match else text


def _month_year(value: Any) -> tuple[str, str, str]:
    text = "" if value is None else str(value).strip()
    if not text or text.lower() in {"none", "null"}:
        return "", "", ""
    if text.lower() == "present":
        return "present", "present", ""
    if re.fullmatch(r"\d{4}", text):
        return text, "", text
    month = re.fullmatch(r"(\d{4})-(\d{2})", text)
    if month:
        year, mm = month.group(1), int(month.group(2))
        if 1 <= mm <= 12:
            return f"{_MONTHS[mm - 1]} {year}", _MONTHS[mm - 1], year
    return text, text, ""


def _ats_date_range(entry: dict[str, Any]) -> str:
    explicit = entry.get("date")
    if isinstance(explicit, str) and explicit.strip():
        return explicit.strip()
    start_disp, start_m, start_y = _month_year(entry.get("start_date"))
    end_disp, end_m, end_y = _month_year(entry.get("end_date"))
    if start_m and end_m and start_y and start_y == end_y:
        return f"{start_m} to {end_m} {end_y}"
    if start_disp and end_disp:
        return f"{start_disp} to {end_disp}"
    return start_disp or end_disp


def _contact(cv: dict[str, Any]) -> dict[str, str]:
    github = "https://github.com/UdayRaj2003"
    linkedin = "https://www.linkedin.com/in/uday-raj-gupta-b7a493262/"
    gfg = "https://www.geeksforgeeks.org/profile/udayrajgusc9z"
    leetcode = "https://leetcode.com/u/udayrajgupta2003/"
    for custom in cv.get("custom_connections") or []:
        if not isinstance(custom, dict):
            continue
        label = str(custom.get("placeholder") or "").lower()
        url = str(custom.get("url") or "").strip()
        if "git" in label:
            github = url or github
        elif "linked" in label:
            linkedin = url or linkedin
        elif "gfg" in label or "geeks" in label:
            gfg = url or gfg
        elif "leet" in label:
            leetcode = url or leetcode
    return {
        "name": str(cv.get("name") or "").strip(),
        "email": str(cv.get("email") or "").strip(),
        "phone": str(cv.get("phone") or "").strip(),
        "location": str(cv.get("location") or "").strip(),
        "github": github,
        "linkedin": linkedin,
        "gfg": gfg,
        "leetcode": leetcode,
    }


def _href_for(chunk: str) -> str:
    cleaned = chunk.rstrip(".,);")
    if cleaned.lower().startswith(("http://", "https://", "mailto:", "tel:")):
        return cleaned
    return "https://" + cleaned


def _url_segments(text: str, explicit_url: str | None = None) -> list[tuple[str, str | None]]:
    """Split text so URL-looking spans can become clickable annotations."""
    if explicit_url:
        return [(text, explicit_url)]
    parts: list[tuple[str, str | None]] = []
    pos = 0
    for match in _URL_CHUNK.finditer(text):
        if match.start() > pos:
            parts.append((text[pos : match.start()], None))
        chunk = match.group(0)
        parts.append((chunk, _href_for(chunk)))
        pos = match.end()
    if pos < len(text):
        parts.append((text[pos:], None))
    return parts or [(text, None)]


def _display_url(url: str) -> str:
    return re.sub(r"^https?://(www\.)?", "", url).rstrip("/")


def _labeled_runs(text: str) -> list[AtsRun]:
    raw = _plain(str(text or "")).strip()
    match = _LABEL.match(raw)
    if match:
        return [
            AtsRun(f"{match.group(3).strip()}: ", bold=True, size=SIZE_BODY),
            AtsRun(match.group(5).strip(), size=SIZE_BODY),
        ]
    return [AtsRun(raw, size=SIZE_BODY)]


def ats_paragraphs(cv: dict[str, Any]) -> list[AtsPara]:
    c = _contact(cv)
    paras: list[AtsPara] = []

    def add(runs: list[AtsRun], indent_cm: float = 0.0) -> None:
        paras.append(AtsPara(runs=runs, indent_cm=indent_cm))

    def kv(label: str, value: str, url: str | None = None, comma: bool = True) -> list[AtsRun]:
        runs = [
            AtsRun(f"{label} : ", bold=True, size=SIZE_LABEL),
            AtsRun(value, size=SIZE_LABEL, url=url),
        ]
        if comma:
            runs.append(AtsRun(" , ", size=SIZE_LABEL))
        return runs

    header: list[AtsRun] = []
    header += kv("NAME", c["name"])
    header += kv("Email", c["email"], f"mailto:{c['email']}" if c["email"] else None)
    tel = re.sub(r"[^\d+]", "", c["phone"])
    header += kv("Phone", c["phone"], f"tel:{tel}" if tel else None)
    header += kv("Location", c["location"])
    header += kv("LinkedIn", _display_url(c["linkedin"]), c["linkedin"])
    header += kv("GitHub", _display_url(c["github"]), c["github"], comma=False)
    add(header)
    add(
        kv("GFG", _display_url(c["gfg"]), c["gfg"])
        + kv("LeetCode", _display_url(c["leetcode"]), c["leetcode"], comma=False)
    )

    sections = cv.get("sections") or {}
    if not isinstance(sections, dict):
        return paras

    for section_name, entries in sections.items():
        if not entries:
            continue
        items = list(entries) if isinstance(entries, list) else [entries]
        key = str(section_name).strip().lower()

        if key in {"profile", "professional summary"}:
            first = items[0]
            text = first if isinstance(first, str) else str((first or {}).get("bullet") or "")
            add(
                [AtsRun("PROFESSIONAL SUMMARY : ", bold=True, size=SIZE_LABEL)]
                + [AtsRun(_plain(text).strip(), size=SIZE_BODY)]
            )
            continue
        if key == "skills":
            add([AtsRun("SKILLS :", bold=True, size=SIZE_LABEL)])
            for entry in items:
                if isinstance(entry, str):
                    add(_labeled_runs(entry))
                elif isinstance(entry, dict):
                    label = str(entry.get("label") or "").strip()
                    details = str(entry.get("details") or "").strip()
                    add(_labeled_runs(f"{label}: {details}" if label else details))
            continue
        if key == "experience":
            add([AtsRun("EXPERIENCE :", bold=True, size=SIZE_LABEL)])
            for entry in items:
                if not isinstance(entry, dict):
                    continue
                company = _strip_link(str(entry.get("company") or ""))
                loc = str(entry.get("location") or "").strip()
                position = str(entry.get("position") or "").strip()
                line = f"{position} , COMPANY : {company}"
                if loc:
                    line += f", {loc}"
                line += f" , DATE : {_ats_date_range(entry)}"
                add([AtsRun(line, bold=True, size=SIZE_LABEL)])
                for hl in entry.get("highlights") or []:
                    if hl:
                        add(_labeled_runs(str(hl)))
            continue
        if key == "projects":
            add([AtsRun("PROJECTS:", bold=True, size=SIZE_LABEL)])
            for entry in items:
                if not isinstance(entry, dict):
                    continue
                title = _strip_link(str(entry.get("name") or ""))
                role = str(entry.get("summary") or "").strip()
                add(
                    [
                        AtsRun(
                            f"TITLE : {title} , ROLE : {role} , DATES : {_ats_date_range(entry)}",
                            bold=True,
                            size=SIZE_LABEL,
                        )
                    ]
                )
                for hl in entry.get("highlights") or []:
                    if hl:
                        add(_labeled_runs(str(hl)))
            continue
        if key == "education":
            for i, entry in enumerate(items):
                if not isinstance(entry, dict):
                    continue
                degree = str(entry.get("degree") or "").strip()
                inst = str(entry.get("institution") or "").strip()
                loc = str(entry.get("location") or "").strip()
                area = str(entry.get("area") or "").strip()
                end_y = _month_year(entry.get("end_date"))[2] or _ats_date_range(entry)
                head = (
                    [AtsRun("EDUCATION : ", bold=True, size=SIZE_LABEL)]
                    if i == 0
                    else [AtsRun(" :  ", bold=True, size=SIZE_BODY)]
                )
                head += [
                    AtsRun("DEGREE: ", bold=True, size=SIZE_BODY),
                    AtsRun(degree, size=SIZE_BODY),
                ]
                add(head)
                inst_line = inst
                if loc:
                    inst_line += f", {loc}"
                if end_y:
                    inst_line += f", Graduated at {end_y}"
                if area:
                    inst_line += f" , {area}"
                add(
                    [
                        AtsRun(" :  INSTITUTION: ", bold=True, size=SIZE_BODY),
                        AtsRun(inst_line, size=SIZE_BODY),
                    ],
                    indent_cm=1.27,
                )
            continue
        if key == "achievements":
            for i, entry in enumerate(items):
                text = entry if isinstance(entry, str) else str((entry or {}).get("bullet") or "")
                if i == 0:
                    add(
                        [AtsRun("ACHIEVEMENTS  : ", bold=True, size=SIZE_LABEL)]
                        + _labeled_runs(text)
                    )
                else:
                    add(
                        [AtsRun(" :  ", bold=True, size=SIZE_BODY)] + _labeled_runs(text),
                        indent_cm=1.27,
                    )
            continue
        if key == "certifications":
            add([AtsRun("CERTIFICATIONS", bold=True, size=SIZE_LABEL)])
            for entry in items:
                text = entry if isinstance(entry, str) else str((entry or {}).get("bullet") or "")
                add(_labeled_runs(text))
            continue
        add([AtsRun(f"{str(section_name).upper()} :", bold=True, size=SIZE_LABEL)])
        for entry in items:
            if isinstance(entry, str):
                add(_labeled_runs(entry))
            elif isinstance(entry, dict):
                add(_labeled_runs(str(entry.get("bullet") or entry.get("details") or "")))
    return paras


def _load_cv(yaml_path: Path) -> dict[str, Any] | None:
    yaml = YAML(typ="rt")
    with Path(yaml_path).open("r", encoding="utf-8") as fh:
        data = yaml.load(fh)
    if not isinstance(data, dict):
        return None
    cv = data.get("cv")
    return cv if isinstance(cv, dict) else None


def _set_run_font(run, *, name: str, size_pt: float, bold: bool = False, color=None) -> None:
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size_pt)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _add_hyperlink(paragraph, text: str, url: str, *, font_name: str, size_pt: float, bold: bool = False) -> None:
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    r_pr.append(r_fonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size_pt * 2)))
    r_pr.append(sz)
    if bold:
        r_pr.append(OxmlElement("w:b"))
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), LINK_BLUE_HEX)
    r_pr.append(color_el)
    new_run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.set(qn("xml:space"), "preserve")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def _tight(paragraph) -> None:
    pf = paragraph.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE


def render_yaml_to_docx(yaml_path: Path, docx_path: Path) -> Path | None:
    if Document is None:
        print("python-docx is not installed; skipping formatted DOCX.", file=sys.stderr, flush=True)
        return None
    cv = _load_cv(yaml_path)
    if cv is None:
        return None

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(DOCX_TOP_IN)
    section.bottom_margin = Inches(DOCX_EDGE_IN)
    section.left_margin = Inches(DOCX_EDGE_IN)
    section.right_margin = Inches(DOCX_EDGE_IN)
    section.header_distance = Inches(0)
    section.footer_distance = Inches(0)

    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(SIZE_BODY)
    if BLACK is not None:
        style.font.color.rgb = BLACK
    nfp = style.paragraph_format
    nfp.space_before = Pt(0)
    nfp.space_after = Pt(0)
    nfp.line_spacing = 1.0

    for block in ats_paragraphs(cv):
        p = doc.add_paragraph()
        _tight(p)
        if block.indent_cm:
            p.paragraph_format.left_indent = Cm(block.indent_cm)
        for run in block.runs:
            for piece, href in _url_segments(run.text, run.url):
                if href:
                    _add_hyperlink(
                        p,
                        piece,
                        href,
                        font_name=FONT,
                        size_pt=run.size,
                        bold=run.bold,
                    )
                else:
                    wr = p.add_run(piece)
                    _set_run_font(wr, name=FONT, size_pt=run.size, bold=run.bold, color=BLACK)

    out = Path(docx_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    if not out.is_file() or out.stat().st_size <= 0:
        return None
    return out


def _calibri_fonts() -> tuple[Path | None, Path | None]:
    fonts = Path(r"C:\Windows\Fonts")
    regular = fonts / "calibri.ttf"
    bold = fonts / "calibrib.ttf"
    return (
        regular if regular.is_file() else None,
        bold if bold.is_file() else None,
    )


def render_yaml_to_pdf_ats(yaml_path: Path, pdf_path: Path) -> Path | None:
    """PDF with the same labeled ATS layout as the DOCX (not RenderCV classic)."""
    try:
        from fpdf import FPDF
    except ImportError:
        print("fpdf2 is not installed; keeping RenderCV PDF.", file=sys.stderr, flush=True)
        return None

    cv = _load_cv(yaml_path)
    if cv is None:
        return None

    regular, bold = _calibri_fonts()
    pdf = FPDF(format="Letter", unit="mm")
    # 1.3mm ≈ 0.05in top
    pdf.set_margins(left=3.0, top=1.3, right=3.0)
    pdf.set_auto_page_break(auto=True, margin=4.0)
    pdf.add_page()
    family = "Helvetica"
    if regular is not None:
        pdf.add_font("Calibri", "", str(regular))
        family = "Calibri"
    if bold is not None:
        pdf.add_font("Calibri", "B", str(bold))
    pdf.set_text_color(0, 0, 0)

    left = pdf.l_margin
    for block in ats_paragraphs(cv):
        indent = block.indent_cm * 10.0
        pdf.set_left_margin(left + indent)
        pdf.set_x(left + indent)
        line_h = 4.2
        for run in block.runs:
            style = "B" if run.bold else ""
            try:
                pdf.set_font(family, style, run.size)
            except Exception:
                pdf.set_font("Helvetica", style, run.size)
            line_h = 4.4 if run.size >= 12 else 3.7
            clean_text = str(run.text or "").replace("\u2011", "-")
            for piece, href in _url_segments(clean_text, run.url):
                if href:
                    pdf.set_text_color(*LINK_BLUE_RGB)
                    pdf.write(line_h, piece, link=href)
                    pdf.set_text_color(0, 0, 0)
                else:
                    pdf.set_text_color(0, 0, 0)
                    pdf.write(line_h, piece)
        pdf.ln(line_h)
        pdf.set_left_margin(left)

    out = Path(pdf_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(out))
    if not out.is_file() or out.stat().st_size <= 0:
        return None
    return out
